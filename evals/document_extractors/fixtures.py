"""Hermetic fixture generation for document extractor evaluations."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

_GREEK_ALPHA = "α"  # noqa: RUF001 - intentional script coverage.
_RTL_TEXT = "العربية ثم עברית ثم 中文"
_PDF_ENCRYPT_AES_256 = 5  # PyMuPDF's public constant is absent from its type information.


@dataclass(frozen=True, slots=True)
class ExpectedSegment:
    """Exact structure expected from a generated fixture."""

    text: str
    heading: str = ""
    page: int | None = None


@dataclass(frozen=True, slots=True)
class Fixture:
    """One evaluation input and its manually defined smoke expectations."""

    fixture_id: str
    path: Path
    valid: bool
    selection_scope: bool = True
    expected_text: str = ""
    anchors: tuple[str, ...] = ()
    table_rows: tuple[str, ...] = ()
    expected_segments: tuple[ExpectedSegment, ...] = ()
    expected_markdown_headings: tuple[tuple[int, str], ...] = ()


def create_docx_smoke_fixtures(directory: Path) -> list[Fixture]:
    """Create deterministic DOCX smoke fixtures in ``directory``."""
    directory.mkdir(parents=True, exist_ok=True)

    structured_path = directory / "structured.docx"
    document = Document()
    document.add_heading("总体报告", level=0)
    document.add_paragraph("前言")
    document.add_heading("第三章", level=1)
    document.add_paragraph("第三章正文")
    document.add_heading("高温耐久性", level=2)
    document.add_paragraph("详细内容")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "样本"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "1.0"
    document.add_paragraph("结尾段")
    document.save(str(structured_path))

    split_runs_path = directory / "split-runs.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Rust ").bold = True
    paragraph.add_run("抽取器")
    document.add_paragraph("English and 中文 stay in reading order.")
    document.add_paragraph("")
    document.add_paragraph("最后一段")
    document.save(str(split_runs_path))

    corrupt_path = directory / "corrupt.docx"
    corrupt_path.write_bytes(b"this is not an OOXML package")

    return [
        Fixture(
            fixture_id="structured",
            path=structured_path,
            valid=True,
            expected_text=("总体报告\n前言\n第三章\n第三章正文\n高温耐久性\n详细内容\n样本 | 数值\nA | 1.0\n结尾段"),
            anchors=(
                "总体报告",
                "前言",
                "第三章",
                "第三章正文",
                "高温耐久性",
                "详细内容",
                "样本",
                "数值",
                "A",
                "1.0",
                "结尾段",
            ),
            table_rows=("样本 | 数值", "A | 1.0"),
            expected_segments=(
                ExpectedSegment("总体报告\n前言", "总体报告"),
                ExpectedSegment("第三章\n第三章正文", "总体报告 / 第三章"),
                ExpectedSegment(
                    "高温耐久性\n详细内容\n样本 | 数值\nA | 1.0\n结尾段",
                    "总体报告 / 第三章 / 高温耐久性",
                ),
            ),
            expected_markdown_headings=((1, "总体报告"), (2, "第三章"), (3, "高温耐久性")),
        ),
        Fixture(
            fixture_id="split-runs",
            path=split_runs_path,
            valid=True,
            expected_text="Rust 抽取器\nEnglish and 中文 stay in reading order.\n最后一段",
            anchors=("Rust 抽取器", "English and 中文 stay in reading order.", "最后一段"),
            expected_segments=(ExpectedSegment("Rust 抽取器\nEnglish and 中文 stay in reading order.\n最后一段"),),
            expected_markdown_headings=(),
        ),
        Fixture(fixture_id="corrupt", path=corrupt_path, valid=False),
    ]


def create_docx_quality_fixtures(directory: Path) -> list[Fixture]:
    """Create the generated DOCX quality corpus with exact structural truth."""
    fixtures = create_docx_smoke_fixtures(directory)

    hierarchy_path = directory / "heading-hierarchy.docx"
    document = Document()
    document.add_paragraph("无标题前导")
    document.add_heading("项目 Atlas", level=0)
    document.add_paragraph("概览正文")
    document.add_heading("方法", level=1)
    document.add_paragraph("方法正文")
    document.add_heading("跳级细节", level=3)
    document.add_paragraph("跳级正文")
    document.add_heading("二级结论", level=2)
    document.add_paragraph("结论正文")
    document.add_heading("结果", level=1)
    document.add_paragraph("结果正文")
    document.save(str(hierarchy_path))

    interleaved_path = directory / "interleaved-tables.docx"
    document = Document()
    document.add_heading("数据", level=1)
    document.add_paragraph("表格之前")
    first_table = document.add_table(rows=2, cols=3)
    for row, values in zip(
        first_table.rows,
        (("名称", "状态", "计数"), (_GREEK_ALPHA, "正常", "2")),
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("两个表格之间")
    second_table = document.add_table(rows=1, cols=2)
    second_table.cell(0, 0).text = "合计"
    second_table.cell(0, 1).text = "2"
    document.add_paragraph("表格之后")
    document.save(str(interleaved_path))

    text_features_path = directory / "text-features.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("前缀")
    paragraph.add_run("与拆分运行")
    paragraph.add_run().add_tab()
    paragraph.add_run("制表后")
    document.add_paragraph(_RTL_TEXT)
    styles = document.styles
    # python-docx omits this method's annotations, but its runtime API is stable.
    styles.add_style("NotAHeading", WD_STYLE_TYPE.PARAGRAPH)  # pyright: ignore[reportUnknownMemberType]
    document.add_paragraph("自定义样式不是标题", style="NotAHeading")
    document.save(str(text_features_path))

    fixtures.extend(
        [
            Fixture(
                fixture_id="heading-hierarchy",
                path=hierarchy_path,
                valid=True,
                expected_text=(
                    "无标题前导\n项目 Atlas\n概览正文\n方法\n方法正文\n跳级细节\n跳级正文\n"
                    "二级结论\n结论正文\n结果\n结果正文"
                ),
                anchors=("无标题前导", "项目 Atlas", "方法", "跳级细节", "二级结论", "结果", "结果正文"),
                expected_segments=(
                    ExpectedSegment("无标题前导"),
                    ExpectedSegment("项目 Atlas\n概览正文", "项目 Atlas"),
                    ExpectedSegment("方法\n方法正文", "项目 Atlas / 方法"),
                    ExpectedSegment("跳级细节\n跳级正文", "项目 Atlas / 方法 / 跳级细节"),
                    ExpectedSegment("二级结论\n结论正文", "项目 Atlas / 方法 / 二级结论"),
                    ExpectedSegment("结果\n结果正文", "项目 Atlas / 结果"),
                ),
                expected_markdown_headings=(
                    (1, "项目 Atlas"),
                    (2, "方法"),
                    (4, "跳级细节"),
                    (3, "二级结论"),
                    (2, "结果"),
                ),
            ),
            Fixture(
                fixture_id="interleaved-tables",
                path=interleaved_path,
                valid=True,
                expected_text=(
                    f"数据\n表格之前\n名称 | 状态 | 计数\n{_GREEK_ALPHA} | 正常 | 2\n两个表格之间\n合计 | 2\n表格之后"
                ),
                anchors=("数据", "表格之前", "名称", "正常", "两个表格之间", "合计", "表格之后"),
                table_rows=("名称 | 状态 | 计数", f"{_GREEK_ALPHA} | 正常 | 2", "合计 | 2"),
                expected_segments=(
                    ExpectedSegment(
                        f"数据\n表格之前\n名称 | 状态 | 计数\n{_GREEK_ALPHA} | 正常 | 2\n"
                        "两个表格之间\n合计 | 2\n表格之后",
                        "数据",
                    ),
                ),
                expected_markdown_headings=((1, "数据"),),
            ),
            Fixture(
                fixture_id="text-features",
                path=text_features_path,
                valid=True,
                expected_text=f"前缀与拆分运行\t制表后\n{_RTL_TEXT}\n自定义样式不是标题",
                anchors=("前缀与拆分运行", "制表后", "中文", "自定义样式不是标题"),
                expected_segments=(ExpectedSegment(f"前缀与拆分运行\t制表后\n{_RTL_TEXT}\n自定义样式不是标题"),),
                expected_markdown_headings=(),
            ),
        ]
    )
    return fixtures


def create_pdf_smoke_fixtures(directory: Path) -> list[Fixture]:
    """Create deterministic PDF fixtures for API and safety checks."""
    directory.mkdir(parents=True, exist_ok=True)

    outline_path = directory / "multipage-outline.pdf"
    document = pymupdf.open()
    page_text = (
        ("Atlas report", "Introduction"),
        ("Methods", "Step one"),
        ("Continued method",),
        (),
        ("Results", "Complete"),
    )
    for lines in page_text:
        page = document.new_page()
        for line_number, line in enumerate(lines):
            page.insert_text(  # pyright: ignore[reportUnknownMemberType]
                (72, 72 + line_number * 24), line, fontsize=20 if line_number == 0 else 11
            )
    document.set_toc(  # pyright: ignore[reportUnknownMemberType]
        [[1, "Atlas report", 1], [2, "Methods", 2], [2, "Results", 5]]
    )
    document.save(outline_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    columns_path = directory / "two-columns.pdf"
    document = pymupdf.open()
    page = document.new_page()
    for line_number, line in enumerate(("Left one", "Left two", "Left three")):
        page.insert_text((72, 72 + line_number * 20), line)  # pyright: ignore[reportUnknownMemberType]
    for line_number, line in enumerate(("Right one", "Right two", "Right three")):
        page.insert_text((300, 72 + line_number * 20), line)  # pyright: ignore[reportUnknownMemberType]
    document.save(columns_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    rotated_path = directory / "rotated.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Rotated page text")  # pyright: ignore[reportUnknownMemberType]
    page.set_rotation(90)  # pyright: ignore[reportUnknownMemberType]
    document.save(rotated_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    image_only_path = directory / "image-only.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.draw_rect(  # pyright: ignore[reportUnknownMemberType]
        pymupdf.Rect(72, 72, 200, 200), fill=(0.2, 0.4, 0.8)
    )
    document.save(image_only_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    encrypted_path = directory / "encrypted.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Secret text")  # pyright: ignore[reportUnknownMemberType]
    document.save(  # pyright: ignore[reportUnknownMemberType]
        encrypted_path,
        encryption=_PDF_ENCRYPT_AES_256,
        owner_pw="owner-password",
        user_pw="user-password",
    )
    document.close()

    corrupt_path = directory / "corrupt.pdf"
    corrupt_path.write_bytes(b"%PDF-1.7\nthis is not a valid PDF")

    return [
        Fixture(
            fixture_id="multipage-outline",
            path=outline_path,
            valid=True,
            expected_text="Atlas report\nIntroduction\nMethods\nStep one\nContinued method\nResults\nComplete",
            anchors=("Atlas report", "Introduction", "Methods", "Step one", "Continued method", "Results", "Complete"),
            expected_segments=(
                ExpectedSegment("Atlas report\nIntroduction", "Atlas report", 1),
                ExpectedSegment("Methods\nStep one", "Atlas report / Methods", 2),
                ExpectedSegment("Continued method", "Atlas report / Methods", 3),
                ExpectedSegment("Results\nComplete", "Atlas report / Results", 5),
            ),
        ),
        Fixture(
            fixture_id="two-columns",
            path=columns_path,
            valid=True,
            selection_scope=False,
            expected_text="Left one\nLeft two\nLeft three\nRight one\nRight two\nRight three",
            anchors=("Left one", "Left two", "Left three", "Right one", "Right two", "Right three"),
            expected_segments=(
                ExpectedSegment("Left one\nLeft two\nLeft three\nRight one\nRight two\nRight three", page=1),
            ),
        ),
        Fixture(
            fixture_id="rotated",
            path=rotated_path,
            valid=True,
            expected_text="Rotated page text",
            anchors=("Rotated page text",),
            expected_segments=(ExpectedSegment("Rotated page text", page=1),),
        ),
        Fixture(
            fixture_id="image-only",
            path=image_only_path,
            valid=True,
            expected_segments=(),
        ),
        Fixture(fixture_id="encrypted", path=encrypted_path, valid=False),
        Fixture(fixture_id="corrupt", path=corrupt_path, valid=False),
    ]


def create_pdf_quality_fixtures(directory: Path) -> list[Fixture]:
    """Create PDF quality fixtures within the supported single-flow scope."""
    fixtures = create_pdf_smoke_fixtures(directory)

    formatting_path = directory / "single-column-formatting.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Quality Report", fontsize=24)  # pyright: ignore[reportUnknownMemberType]
    page.insert_text((72, 112), "Overview", fontsize=16)  # pyright: ignore[reportUnknownMemberType]
    page.insert_text((72, 142), "First paragraph in reading order.")  # pyright: ignore[reportUnknownMemberType]
    page.insert_text((72, 166), "Second paragraph remains complete.")  # pyright: ignore[reportUnknownMemberType]
    document.save(formatting_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    cjk_path = directory / "cjk.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text(  # pyright: ignore[reportUnknownMemberType]
        (72, 72), "中文抽取保持顺序。第二句完整。", fontname="china-s", fontsize=14
    )
    document.save(cjk_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    header_footer_path = directory / "header-footer.pdf"
    document = pymupdf.open()
    header_segments: list[ExpectedSegment] = []
    header_texts: list[str] = []
    for page_number in range(1, 4):
        page = document.new_page()
        lines = ("Quarterly report", f"Body section {page_number}", f"Page {page_number}")
        page.insert_text((72, 42), lines[0], fontsize=9)  # pyright: ignore[reportUnknownMemberType]
        page.insert_text((72, 100), lines[1], fontsize=12)  # pyright: ignore[reportUnknownMemberType]
        page.insert_text((72, 790), lines[2], fontsize=9)  # pyright: ignore[reportUnknownMemberType]
        text = "\n".join(lines)
        header_texts.append(text)
        header_segments.append(ExpectedSegment(text, page=page_number))
    document.save(header_footer_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    incremental_path = directory / "incremental.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Original revision")  # pyright: ignore[reportUnknownMemberType]
    document.save(incremental_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()
    document = pymupdf.open(incremental_path)
    page = document.new_page()
    page.insert_text((72, 72), "Incremental revision")  # pyright: ignore[reportUnknownMemberType]
    document.saveIncr()  # pyright: ignore[reportUnknownMemberType]
    document.close()

    form_path = directory / "filled-form.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Name:")  # pyright: ignore[reportUnknownMemberType]
    widget = pymupdf.Widget()
    widget.field_name = "name"
    widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    widget.field_value = "Alice Example"
    widget.rect = pymupdf.Rect(120, 55, 300, 80)
    page.add_widget(widget)  # pyright: ignore[reportUnknownMemberType]
    document.save(form_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    literals_path = directory / "markdown-literals.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Markdown Safety", fontsize=22)  # pyright: ignore[reportUnknownMemberType]
    page.insert_text((72, 112), "Literal * stars * and `backticks` stay readable.")  # pyright: ignore[reportUnknownMemberType]
    page.insert_text((72, 136), "underscore_name and ~~markers~~ are content.")  # pyright: ignore[reportUnknownMemberType]
    document.save(literals_path)  # pyright: ignore[reportUnknownMemberType]
    document.close()

    fixtures.extend(
        [
            Fixture(
                fixture_id="single-column-formatting",
                path=formatting_path,
                valid=True,
                expected_text=(
                    "Quality Report\nOverview\nFirst paragraph in reading order.\nSecond paragraph remains complete."
                ),
                anchors=("Quality Report", "Overview", "First paragraph", "Second paragraph"),
                expected_segments=(
                    ExpectedSegment(
                        "Quality Report\nOverview\nFirst paragraph in reading order.\n"
                        "Second paragraph remains complete.",
                        page=1,
                    ),
                ),
                expected_markdown_headings=((1, "Quality Report"), (2, "Overview")),
            ),
            Fixture(
                fixture_id="cjk",
                path=cjk_path,
                valid=True,
                expected_text="中文抽取保持顺序。第二句完整。",
                anchors=("中文抽取", "第二句完整"),
                expected_segments=(ExpectedSegment("中文抽取保持顺序。第二句完整。", page=1),),
            ),
            Fixture(
                fixture_id="header-footer",
                path=header_footer_path,
                valid=True,
                expected_text="\n".join(header_texts),
                anchors=("Body section 1", "Body section 2", "Body section 3"),
                expected_segments=tuple(header_segments),
            ),
            Fixture(
                fixture_id="incremental",
                path=incremental_path,
                valid=True,
                expected_text="Original revision\nIncremental revision",
                anchors=("Original revision", "Incremental revision"),
                expected_segments=(
                    ExpectedSegment("Original revision", page=1),
                    ExpectedSegment("Incremental revision", page=2),
                ),
            ),
            Fixture(
                fixture_id="filled-form",
                path=form_path,
                valid=True,
                expected_text="Name:\nAlice Example",
                anchors=("Name:", "Alice Example"),
                expected_segments=(ExpectedSegment("Name:\nAlice Example", page=1),),
            ),
            Fixture(
                fixture_id="markdown-literals",
                path=literals_path,
                valid=True,
                expected_text=(
                    "Markdown Safety\nLiteral * stars * and `backticks` stay readable.\n"
                    "underscore_name and ~~markers~~ are content."
                ),
                anchors=("Markdown Safety", "Literal * stars *", "underscore_name", "~~markers~~"),
                expected_segments=(
                    ExpectedSegment(
                        "Markdown Safety\nLiteral * stars * and `backticks` stay readable.\n"
                        "underscore_name and ~~markers~~ are content.",
                        page=1,
                    ),
                ),
                expected_markdown_headings=((1, "Markdown Safety"),),
            ),
        ]
    )
    return fixtures
